"""Build the illustrated ICDS/CAQ group-meeting deck from authoritative files.

Charts read experiment JSON; diagrams/text are editable PowerPoint objects.
Run with the existing tabicl Python. No network, training, or data regeneration.
"""
import io
import json
import math
import zipfile
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

EXP = Path(__file__).resolve().parents[1]
ASSETS = Path(__file__).resolve().parent / "assets"
ASSETS.mkdir(parents=True, exist_ok=True)
ICDS = Path("/home/lizitao/project/Industrial_synthetic_data")
DOC = ICDS / "reports/icds-method-report/ICDS_工业数据合成方法汇报.docx"
DATA = ICDS / "data/industrial-v2"
RESULTS = EXP / "checkpoints/industrial-v2"
TEST = RESULTS / "meta-test-seed42-context960-query256-official8"
summary = json.loads((DATA / "dataset-summary.json").read_text())
validation = json.loads((RESULTS / "formal-comparison.json").read_text())
test = json.loads((TEST / "report.json").read_text())
assert test["status"] == "completed" and validation["status"] == "verified_completed"
assert test["checkpoints"]["caq"]["selected_step"] == 19000
assert test["checkpoints"]["tabicl_ft"]["selected_step"] == 15000

FONT = "Noto Sans CJK SC"
C = {"navy": "163555", "teal": "008A8D", "blue": "5279A8", "orange": "CE7631",
     "purple": "7C64A8", "ink": "20354B", "muted": "65788C", "light": "EDF4F7",
     "border": "D8E3EA", "white": "FFFFFF", "green": "168B76", "red": "BE6A46"}
plt.rcParams.update({"font.family": FONT, "axes.unicode_minus": False,
                     "font.size": 15, "axes.spines.top": False, "axes.spines.right": False,
                     "axes.labelcolor": "#"+C["ink"], "text.color": "#"+C["ink"]})
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
prs.core_properties.title = "ICDS工业合成数据与CAQ-TabICL预测实验"
prs.core_properties.subject = "研究生组会｜方法组成、8类原型、industrial-v2实验"
prs.core_properties.author = "CAQ-TabICL / ICDS研究项目"
prs.core_properties.keywords = "industrial-v2, ICDS, CAQ-TabICL, seed42, meta_test"
outline = []


def rgb(color):
    return RGBColor.from_string(C.get(color, color))


def box(s, x, y, w, h, fill="light", line=None, rounded=True):
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                              Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shape.adjustments[0] = .12
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line); shape.line.width = Pt(1.3)
    else:
        shape.line.fill.background()
    return shape


def text(s, value, x, y, w, h, size=22, color="ink", bold=False, align=PP_ALIGN.LEFT,
         valign=MSO_ANCHOR.TOP):
    shape = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.02)
    frame.margin_top = frame.margin_bottom = Inches(.01)
    frame.vertical_anchor = valign
    for i, line in enumerate(str(value).split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line; p.alignment = align
        p.space_after = Pt(5)
        for run in p.runs:
            run.font.name = FONT; run.font.size = Pt(size)
            run.font.bold = bold; run.font.color.rgb = rgb(color)
            ea = OxmlElement("a:ea"); ea.set("typeface", FONT)
            run._r.get_or_add_rPr().append(ea)
    return shape


def line(s, x1, y1, x2, y2, color="teal", width=2.1, arrow=False):
    shape = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color); shape.line.width = Pt(width)
    if arrow:
        end = OxmlElement("a:tailEnd"); end.set("type", "triangle")
        shape.line._get_or_add_ln().append(end)
    return shape


def dot(s, x, y, d=.28, color="teal"):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(color); sh.line.fill.background()
    return sh


def pill(s, value, x, y, w, color="teal"):
    box(s, x, y, w, .37, color)
    text(s, value, x+.07, y+.025, w-.14, .3, 13, "white", True, PP_ALIGN.CENTER)


def slide(title, subtitle, note, source="稿件：ICDS v0.3.0｜数据：industrial-v2"):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb("white")
    box(s, 0, 0, 13.333, .09, "teal", rounded=False)
    text(s, title, .55, .32, 12.1, .66, 30, "navy", True)
    text(s, subtitle, .58, 1.0, 12.0, .44, 16, "muted")
    text(s, source, .58, 7.07, 11.6, .24, 10, "muted")
    text(s, f"{len(prs.slides):02d}", 12.35, 7.03, .4, .28, 11, "muted", align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = (note + "\n\n核对来源：" + str(DOC) + "\n" +
                                          str(DATA / "dataset-summary.json") + "\n" +
                                          str(RESULTS / "formal-comparison.json") + "\n" +
                                          str(TEST / "report.json"))
    outline.append((title, note))
    return s


def picture(s, path, x, y, w, h):
    im = Image.open(path) if not isinstance(path, io.BytesIO) else Image.open(io.BytesIO(path.getvalue()))
    ratio = im.width/im.height
    pw, ph = min(w, h*ratio), min(h, w/ratio)
    if isinstance(path, io.BytesIO): path.seek(0)
    return s.shapes.add_picture(path if isinstance(path, io.BytesIO) else str(path),
                               Inches(x+(w-pw)/2), Inches(y+(h-ph)/2), width=Inches(pw), height=Inches(ph))


def takeaway(s, value, color="teal"):
    box(s, .58, 6.47, 12.15, .43, "light")
    text(s, value, .77, 6.50, 11.75, .35, 16, color, True)


def stat(s, value, label, x, y, w=2.8, color="teal"):
    box(s, x, y, w, 1.18, "light")
    size = 22 if len(value) > 12 else 25 if len(value) > 8 else 30
    value_shape = text(s, value, x+.17, y+.12, w-.34, .53, size, color, True)
    value_shape.text_frame.word_wrap = False
    text(s, label, x+.17, y+.73, w-.34, .3, 14, "muted")


def icon(s, kind, x, y, color="teal", scale=1.):
    def b(dx, dy, w, h, fill="white", rounded=False):
        return box(s, x+dx*scale, y+dy*scale, w*scale, h*scale, fill, color, rounded)
    def l(ax, ay, bx, by, arrow=False):
        return line(s, x+ax*scale, y+ay*scale, x+bx*scale, y+by*scale, color, 1.9, arrow)
    if kind == "abstract_discrete":
        pts = [(.12,.44),(.77,.10),(1.39,.44),(.77,.83)]
        for a,b_ in [(0,1),(1,2),(0,3),(3,2)]: l(*pts[a],*pts[b_])
        for xx,yy in pts: dot(s,x+xx*scale-.08,y+yy*scale-.08,.18*scale,color)
    elif kind == "thermal":
        b(.12,.14,.5,.65); b(.95,.14,.5,.65)
        text(s,"T₁",x+.20*scale,y+.34*scale,.4*scale,.3*scale,15,color)
        text(s,"T₂",x+1.03*scale,y+.34*scale,.4*scale,.3*scale,15,color)
        l(.64,.32,.92,.32,True); l(.92,.63,.64,.63,True)
    elif kind in ("heated_tank","coupled_tanks"):
        for xx in ([.43] if kind == "heated_tank" else [.10,.63,1.16]):
            b(xx,.10,.43,.73)
            box(s,x+(xx+.015)*scale,y+.52*scale,.40*scale,.29*scale,"D2EAEB",rounded=False)
        l(-.04,.31,.08 if kind == "coupled_tanks" else .40,.31,True)
        if kind == "coupled_tanks": l(.54,.57,.61,.57,True); l(1.07,.57,1.14,.57,True)
        else: l(.88,.70,1.38,.70,True)
    elif kind == "mass_spring":
        l(.06,.14,.06,.85)
        coords=[(.07,.47),(.17,.29),(.29,.64),(.41,.29),(.53,.64),(.65,.29),(.73,.47)]
        for a,b_ in zip(coords,coords[1:]): l(*a,*b_)
        b(.76,.25,.48,.49,rounded=True)
        text(s,"m",x+.88*scale,y+.33*scale,.3*scale,.3*scale,17,color)
        l(1.25,.48,1.60,.48,True); l(.04,.87,1.60,.87)
    elif kind == "dc_motor":
        sh=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x+.37*scale), Inches(y+.08*scale), Inches(.81*scale), Inches(.81*scale))
        sh.fill.solid(); sh.fill.fore_color.rgb=rgb("white"); sh.line.color.rgb=rgb(color); sh.line.width=Pt(2)
        text(s,"M",x+.60*scale,y+.29*scale,.4*scale,.4*scale,21,color,True)
        l(.03,.46,.34,.46,True); l(1.18,.46,1.58,.46)
    elif kind == "heat_exchanger":
        b(.17,.14,1.3,.69,rounded=True)
        l(.30,.30,1.30,.30); l(1.30,.65,.30,.65); l(.30,.48,1.30,.48)
        l(.01,.26,.17,.26,True); l(1.62,.71,1.49,.71,True)
    else:
        for xx in [.10,.65,1.20]: b(xx,.27,.40,.44,rounded=True)
        l(.50,.49,.63,.49,True); l(1.05,.49,1.18,.49,True)


# Charts use raw saved metrics, never hard-coded reconstructed scores.
fig, ax = plt.subplots(figsize=(9.6,4.3), layout="constrained")
for name, color, label in [("caq","teal","CAQ"),("tabicl","blue","TabICL-FT")]:
    es=json.loads((RESULTS/f"{name}-A4-seed42-hr0/evaluations.json").read_text())
    xs=[e["step"]/1000 for e in es]; ys=[e["validation"]["macro_fixed_scale_mse"] for e in es]
    ax.plot(xs,ys,color="#"+C[color],label=label,linewidth=2.3)
    r=validation["runs"][name]
    ax.scatter([r["best_step"]/1000],[r["best_mse"]],color="#"+C[color],s=60,zorder=5)
ax.set_yscale("log"); ax.set_xlabel("训练步数 / 千步"); ax.set_ylabel("验证MSE（对数坐标）")
ax.grid(alpha=.18); ax.legend(frameon=False); fig.savefig(ASSETS/"validation.png",dpi=220); plt.close(fig)

vals=[test["models"][m]["prediction"]["macro_fixed_scale_mse"] for m in ("caq","tabicl_ft","tabicl_official_zero_shot")]
fig,axes=plt.subplots(1,2,figsize=(10.8,4.6),gridspec_kw={"width_ratios":[1.25,1]},layout="constrained")
for ax,n,title in [(axes[0],3,"三模型对比｜对数坐标"),(axes[1],2,"结构收益｜线性坐标")]:
    labels=["CAQ","TabICL-FT","官方未微调"][:n]
    bars=ax.bar(labels,vals[:n],color=["#"+C[k] for k in ("teal","blue","muted")][:n],width=.53)
    if n==3:
        ax.set_yscale("log"); ax.set_ylim(.1,20); ax.set_yticks([.1,1,10],labels=["0.1","1","10"])
    else: ax.set_ylim(0,.53)
    for bar,v in zip(bars,vals[:n]): ax.text(bar.get_x()+bar.get_width()/2,v*1.10 if n==3 else v+.015,f"{v:.4f}",ha="center",fontsize=16,fontweight="bold")
    ax.set_title(title,fontsize=17,pad=18); ax.set_ylabel("参考尺度归一化MSE ↓"); ax.grid(axis="y",alpha=.15); ax.set_axisbelow(True)
fig.savefig(ASSETS/"test_comparison.png",dpi=220); plt.close(fig)

order=["abstract_discrete","thermal","heated_tank","mass_spring","dc_motor","coupled_tanks","heat_exchanger","production_line"]
cn=["抽象系统","热网络","加热液罐","机械执行器","直流电机","耦合液罐","换热器","产线"]
ca=test["models"]["caq"]["prediction"]; ft=test["models"]["tabicl_ft"]["prediction"]
changes=np.array([100*(1-ca["scenario_macro_fixed_scale_mse"][n]/ft["scenario_macro_fixed_scale_mse"][n]) for n in order])
fig,ax=plt.subplots(figsize=(7.6,4.8),layout="constrained")
ax.barh(cn,changes,color=["#"+C["teal" if x>=0 else "orange"] for x in changes],height=.62)
ax.invert_yaxis(); ax.axvline(0,color="#"+C["border"],linewidth=1.3); ax.set_xlim(-29,58)
for i,v in enumerate(changes): ax.text(v+(1.4 if v>=0 else -1.4),i,f"{v:+.1f}%",va="center",ha="left" if v>=0 else "right",fontsize=15,fontweight="bold")
ax.set_xlabel("CAQ相对FT的平均误差降低（%）"); ax.grid(axis="x",alpha=.14); ax.set_axisbelow(True)
fig.savefig(ASSETS/"scenario_reductions.png",dpi=220); plt.close(fig)

fig,ax=plt.subplots(figsize=(4.2,3.8),layout="constrained")
fractions=[summary["coverage"]["families"][f]["rows"] for f in ("physical_ode","hybrid_line","discrete")]
ax.pie(fractions,colors=["#"+C[k] for k in ("teal","purple","blue")],startangle=90,
       wedgeprops={"width":.30,"edgecolor":"white"})
ax.text(0,.09,"183.4万",ha="center",fontsize=27,fontweight="bold"); ax.text(0,-.24,"转移样本行",ha="center",fontsize=14)
fig.savefig(ASSETS/"composition.png",dpi=220); plt.close(fig)

# 01
s=slide("", "", "开场：论文主角是CAQ-TabICL。ICDS提供多机制工业模拟训练任务；本次汇报以新数据及单seed正式测试为证据，不宣称真实工厂泛化。")
box(s,0,0,13.333,7.5,"navy",rounded=False)
pill(s,"研究生组会 · 2026.10.12",.72,.56,3.0,"teal")
text(s,"ICDS工业合成数据\n与CAQ-TabICL预测实验",.72,1.35,11.9,1.95,43,"white",True)
text(s,"多机制合成任务 → 工业训练 → 留出设备的一步预测",.75,3.60,11.8,.55,23,"white")
for i,(v,l) in enumerate([("8类","代表性原型"),("512套","模拟设备"),("183.4万行","一步转移数据")]):
    x=.77+i*4.12; box(s,x,4.83,3.70,1.40,"254B6B")
    text(s,v,x+.22,4.98,3.25,.62,34,"white",True); text(s,l,x+.24,5.80,3.2,.35,17,"CBDCE6")
text(s,"ICDS v0.3.0  /  industrial-v2  /  seed42",.77,6.76,11.8,.33,16,"CBDCE6")

# 02
s=slide("ICDS做什么：把机制变成可训练的统一任务", "不是任意工厂生成器；是有限原型上的共同流程与数据协议", "五个模块统一流程。先选结构再抽参数，后端遵守各自物理或事件机制；同源动作分支为动作响应提供监督。数据、结构和模型收益需分别验证。")
steps=[("场景与\n结构库","变量/单位/连接"),("分层\n随机化","设备/初值/工况"),("动力学\n执行","离散/ODE/事件"),("同源动作\n分支","同起点/换动作"),("质检与\n任务导出","重放/隔离/标签")]
for i,(t,d) in enumerate(steps):
    x=.61+i*2.55; box(s,x,2.04,2.18,2.24,"light","border")
    dot(s,x+.18,2.22,.43,"teal"); text(s,str(i+1),x+.18,2.23,.43,.35,17,"white",True,PP_ALIGN.CENTER)
    text(s,t,x+.17,2.88,1.84,.85,23,"teal",True)
    text(s,d,x+.17,3.78,1.86,.42,14,"muted")
    if i<4: line(s,x+2.21,3.06,x+2.47,3.06,arrow=True)
box(s,.67,4.80,5.65,1.14,"navy"); text(s,"ICDS：训练数据方法",.89,4.97,5.18,.39,24,"white",True); text(s,"回答：数据从哪里来，机制为何合理",.89,5.48,5.16,.32,17,"white")
box(s,6.74,4.80,5.88,1.14,"teal"); text(s,"CAQ-TabICL：论文主角模型",6.96,4.97,5.44,.39,24,"white",True); text(s,"回答：怎样学习动作影响并预测状态",6.96,5.48,5.43,.32,17,"white")
takeaway(s,"统一接口，不统一所有工业过程的具体方程。")

# 03
s=slide("方法原理：统一的一步转移表达", "同一套输入/输出契约，内部可以运行不同的动力学后端", "公式原样引用稿件的image2。x为状态，u为动作，e为当前外部条件，Δt为推进间隔；G为结构，θ为设备参数，ξ为随机创新，b为动作分支。ODE积分与混合事件推进并不是同一个算法。")
with zipfile.ZipFile(DOC) as archive: equation=io.BytesIO(archive.read("word/media/image2.png"))
picture(s,equation,2.05,1.58,9.2,.96)
for i,(t,d) in enumerate([("xₜ：当前状态","温度 / 位置 / 队列"),("uₜ⁽ᵇ⁾：控制动作","加热 / 电压 / 维护"),("eₜ，Δt：条件与间隔","环境驱动 / 推进多久")]):
    y=2.91+i*.85; box(s,.73,y,3.62,.63,"light"); text(s,t,.88,y+.06,3.30,.29,19,"teal",True); text(s,d,.88,y+.37,3.30,.24,12,"muted")
line(s,4.46,4.04,5.05,4.04,arrow=True)
box(s,5.13,2.94,3.49,2.28,"navy"); text(s,"结构 G + 设备参数 θ",5.36,3.18,3.05,.53,21,"white",True)
text(s,"按机制推进一步\n随机创新 ξₜ\n如到货 / 故障",5.36,3.79,3.04,1.25,17,"white")
line(s,8.72,4.04,9.23,4.04,arrow=True)
box(s,9.32,3.39,3.10,1.30,"teal"); text(s,"xₜ₊₁⁽ᵇ⁾",9.59,3.62,2.55,.45,29,"white",True); text(s,"该动作下的下一状态",9.55,4.17,2.59,.36,17,"white")
for i,(t,d) in enumerate([("D：离散更新","实数状态按时间步更新"),("P：ODE积分","变化率 → 数值积分"),("H：混合事件","连续变化 + 加工事件")]):
    x=.79+i*4.16; text(s,t,x,5.59,3.83,.37,20,["blue","teal","purple"][i],True); text(s,d,x,6.02,3.83,.30,15,"muted")

# 04
s=slide("数据层级：图纸 → 设备 → 运行 → 一步及分支", "program、world和样本行数描述不同层次，不能互相替代", "78个program不等于78类工业场景。每个world固定一套设备参数，每world有8个episode，每episode有256个事实步骤。约25%source执行3个替代动作，只有选中的source含四branch。事实轨迹仅沿b0继续。")
levels=[("Program","结构 / 方程 / 连接","78个","电机状态与方程"),("World","一套固定设备参数","512套","电阻 / 惯量 / 热容"),("Episode","一次运行轨迹","4,096条","初值 / 动作 / 外部驱动"),("Source","动作前的一个起点","1,048,576个","同状态测试不同动作")]
for i,(a,b_,v,d) in enumerate(levels):
    x=.64+i*3.19; box(s,x,1.91,2.68,2.62,"light","border")
    text(s,a,x+.17,2.09,2.34,.42,25,"teal",True); text(s,b_,x+.17,2.65,2.34,.68,18)
    text(s,v,x+.17,3.49,2.34,.46,27,"navy",True); text(s,d,x+.17,4.02,2.34,.30,13,"muted")
    if i<3: line(s,x+2.73,3.14,x+3.08,3.14,arrow=True)
text(s,"事实轨迹继续",.78,5.03,2.5,.37,21,"blue",True)
for i in range(3):
    box(s,3.04+i*1.44,4.92,1.06,.53,"blue"); text(s,f"xₜ₊{i}",3.11+i*1.44,4.99,.91,.35,19,"white",True,PP_ALIGN.CENTER)
    if i<2: line(s,4.14+i*1.44,5.19,4.40+i*1.44,5.19,"blue",arrow=True)
text(s,"约25%源点分叉",8.33,4.85,3.85,.35,18,"purple",True)
for i in range(4): pill(s,f"b{i}",8.26+i*.98,5.39,.77,"blue" if i==0 else "purple")
takeaway(s,"共1,834,294行 = 1,048,576事实行 + 785,718替代动作分支行。")

# 05
s=slide("场景组成：1类抽象 + 6类物理 + 1类产线", "8类代表性低阶原型，不是8个工业行业，也不意味着覆盖所有工业过程", "设备数来自dataset-summary.json。机械执行器、直流电机、耦合液罐、多段换热器为本次新增；产线扩展为6种布局。抽象discrete是离散时间，不要求状态取整数。")
descs=["实数状态耦合 / 数学先验","加热 / 散热 / 内部传热","体积 / 热能 / 阀响应","位移 / 速度 / 力滞后","电流 / 转速 / 绕组热","2—4罐串联或分流","1—3段顺流或逆流","2—4工位串联或并联"]
for i,(key,label,desc) in enumerate(zip(order,cn,descs)):
    x=.62+(i%4)*3.17; y=1.76+(i//4)*2.23
    co="blue" if i==0 else "purple" if i==7 else "teal"
    box(s,x,y,2.95,2.02,"light","border"); icon(s,key,x+.55,y+.16,co,.9)
    text(s,label,x+.13,y+1.13,2.68,.39,22,co,True,PP_ALIGN.CENTER)
    text(s,desc,x+.12,y+1.65,2.70,.27,12,"muted",align=PP_ALIGN.CENTER)
    text(s,f"{summary['coverage']['scenarios'][key]['worlds']}套",x+2.12,y+.09,.69,.32,14,co,True,PP_ALIGN.RIGHT)
takeaway(s,"离散时间 ≠ 离散数值；抽象系统同样可以产生连续实数状态。")

# 06
s=slide("产线扩展：2 / 3 / 4工位 × 串联 / 独立并联", "每站6个状态：队列、剩余加工量、忙闲、温度、磨损、故障", "六种有限布局。串联有堵塞和缺料传播；并联为随机分配到独立机台，无共享工人或智能调度。每1秒内10个事件网格，连续温度/磨损与离散到货、故障、维护交织。parallel_2仅4个训练world，不在验证/测试；serial_3无验证world。")
for col,n in enumerate([2,3,4]):
    x=.63+col*4.22
    for row in range(2):
        y=1.78+row*2.05; box(s,x,y,3.95,1.83,"light","border")
        text(s,f"{n}工位{'串联' if row==0 else '独立并联'}",x+.18,y+.13,3.57,.34,21,"purple",True)
        if row==0:
            gap=3.26/n; base=x+.34
            for j in range(n):
                xx=base+j*gap; box(s,xx,y+.78,.56,.57,"purple")
                text(s,str(j+1),xx+.06,y+.89,.42,.27,17,"white",True,PP_ALIGN.CENTER)
                if j<n-1: line(s,xx+.58,y+1.06,base+(j+1)*gap-.07,y+1.06,"purple",arrow=True)
        else:
            gap=3.22/n
            for j in range(n):
                xx=x+.38+j*gap; box(s,xx,y+.87,.57,.54,"purple")
                text(s,str(j+1),xx+.06,y+.97,.43,.29,17,"white",True,PP_ALIGN.CENTER)
                line(s,x+1.98,y+.61,xx+.285,y+.83,"purple",1.2,True)
    
text(s,"控制：速度 + 维护",.83,6.07,3.53,.33,20,"teal",True)
text(s,"串联：堵塞 / 缺料传播",4.96,6.07,3.70,.33,18)
text(s,"并联：机台各自加工",9.13,6.07,3.57,.33,18)
takeaway(s,"布局只是代表性加工机制，不复现完整工厂，也不模拟切削或产品质量。")

# 07: actual arrays, not invented sample values.
ids=np.load(DATA/"shards/3/ids.npy",mmap_mode="r")
ix=np.flatnonzero((ids[:,1]==0)&(ids[:,2]==6147)&(ids[:,4]==3))
ix=ix[np.argsort(ids[ix,3])]; assert len(ix)==4
states=np.load(DATA/"shards/3/state.npy",mmap_mode="r")[ix]
actions=np.load(DATA/"shards/3/action.npy",mmap_mode="r")[ix]
nexts=np.load(DATA/"shards/3/next_state.npy",mmap_mode="r")[ix]
assert np.allclose(states,states[0])
s=slide("同源动作分支：同一个起点，换动作再算", "真实数组样例｜机械执行器 world3 / source6147 / Δt≈0.01043秒", "数值直接读正式shards/3数组。只有动作命令不同；状态、参数和外部条件相同。控制先改变实际力，再影响速度与位移。随机产线分支共享随机创新，但并不自动识别真实系统平均因果效应。")
box(s,.76,2.44,3.55,2.46,"light","border")
text(s,"相同动作前状态",.95,2.62,3.15,.45,24,"teal",True)
text(s,f"位移 {states[0,0]*1000:.4f} mm\n速度 {states[0,1]:.6f} m/s\n实际力 {states[0,2]:.6f} N",.98,3.37,3.02,1.29,20)
text(s,"分支 / 动作命令 u",5.02,1.75,3.11,.43,21,"navy",True)
text(s,"下一实际力 f / N",9.13,1.75,3.18,.43,21,"navy",True)
for i in range(4):
    y=2.44+i*.89; color=["blue","orange","teal","purple"][i]
    line(s,4.38,3.63,4.96,y+.30,color,1.3,True)
    box(s,5.05,y,3.17,.63,color)
    value_shape = text(s,f"b{i}  {actions[i,0]:+.4f}",5.20,y+.10,2.85,.43,21,"white",True)
    value_shape.text_frame.word_wrap = False
    line(s,8.36,y+.31,8.91,y+.31,color,arrow=True)
    text(s,f"{nexts[i,2]:+.6f}",9.23,y+.04,3.02,.51,27,color,True)
takeaway(s,"b0继续原轨迹；b1—b3只记录替代动作的一步结果，不改写事实轨迹。")

# 08
s=slide("数据规模与质量：不只是增加样本行", "原型4 → 8；工业program 7 → 24；产线布局1 → 6（相对旧pilot）", "规模与工业结构丰富度应分别描述。总program从113降至78，是抽象program减少，而工业program变多。8类包含1抽象和7工业原型。P的占比是设计选择，不是工业真实频率。1,834,294行完整验证通过；数值重放和参考积分仅说明实现正确，不等于真实设备标定。", "来源：稿件 §02/06/09；dataset-summary.json")
picture(s,ASSETS/"composition.png",.73,1.81,4.17,3.58)
for i,(label,v,color) in enumerate([("物理ODE","70.52%","teal"),("混合产线","18.91%","purple"),("抽象系统","10.57%","blue")]):
    y=2.10+i*.97; dot(s,5.03,y+.10,.20,color); text(s,label,5.38,y,2.27,.35,20); text(s,v,7.65,y-.04,2.03,.46,27,color,True)
stat(s,"512 / 4,096","设备 / 轨迹",10.03,1.99,2.54)
stat(s,"408 / 52 / 52","训练 / 验证 / 测试设备",10.03,3.49,2.54)
for i,(t,d) in enumerate([("完整验证","183.4万行通过"),("分支重放","抽查2,718个源点"),("参考积分","最大归一化误差6.69×10⁻⁷")]):
    x=.82+i*4.18; text(s,t,x,5.43,3.82,.42,21,"teal",True); text(s,d,x,5.98,3.82,.35,15,"muted")
takeaway(s,"机制合理、数值正确、模型有效：三层证据，不能彼此替代。")

# 09
s=slide("如何训练：context给示例，query让模型作答", "world隔离；同一设备内部再隔离episode与source", "408训练world/52验证/52测试。context来自该world的4个episode，query来自另4个episode，源点及分支不跨context/query。图中的C/Q是分组示意，不指定实际episode编号。连续状态标签为增量，count/binary为下一值。标准化只拟合context。A4的普通MSE覆盖全部query行，并非仅对两两差值算损失。")
box(s,.75,1.85,4.18,3.75,"light","border"); text(s,"一套设备的8条轨迹",.95,2.05,3.78,.43,24,"navy",True)
for j in range(8):
    x=.99+(j%4)*.95; y=2.89+(j//4)*.93; co="blue" if j<4 else "teal"
    box(s,x,y,.78,.60,co); text(s,f"{'C' if j<4 else 'Q'}{j%4+1}",x+.04,y+.13,.70,.30,18,"white",True,PP_ALIGN.CENTER)
text(s,"上：context轨迹\n下：query轨迹",.99,4.85,3.65,.65,17,"muted")
line(s,5.06,3.21,5.54,3.21,arrow=True)
box(s,5.65,2.10,6.96,1.20,"blue"); text(s,"960条context：输入 + 已知答案",5.89,2.35,6.43,.51,26,"white",True)
box(s,5.65,3.63,6.96,1.20,"teal"); text(s,"256条query：输入 → 预测 → 对照真值",5.89,3.86,6.44,.59,24,"white",True)
text(s,"连续量：下一状态 − 当前状态",5.87,5.21,6.42,.38,20)
text(s,"计数 / 二值：下一值；缩放仅由context拟合",5.87,5.72,6.42,.38,18,"muted")
takeaway(s,"训练query配对b0/b1，但普通MSE计算全部query行；b2/b3存储不等于本轮全部训练。")

# 10
s=slide("模型对照：数据收益与CAQ结构收益分开看", "相同官方初始化、相同工业任务序列、各30k步；不是旧工业模型续训", "CAQ和FT冻结ICL predictor、训练列/行模块，CAQ额外训练动作模块。CAQ动作只进专用causal通道，FT用普通特征列。训练24k工业+6k SCM replay，lr1e-4 cosine至5e-6，seed42，full_state hr0，单工业任务1个目标。官方8成员fit/predict只准备上下文，不更新神经网络参数。")
for i,(name,detail,color) in enumerate([("CAQ-TabICL","状态 → 行/列表示\n动作 → 专用causal通道","teal"),("TabICL-FT","状态 + 相同动作\n作为普通特征表输入","blue"),("官方未微调TabICL","fit：960条已知示例\npredict：256条query","muted")]):
    x=.72+i*4.25; box(s,x,1.96,3.97,2.52,"light","border")
    pill(s,"工业训练" if i<2 else "不做梯度训练",x+.19,2.15,2.04,color)
    text(s,name,x+.19,2.86,3.59,.56,25,color,True)
    text(s,detail,x+.21,3.54,3.50,.90,19)
for i,(t,d) in enumerate([("同信息","公开状态 / 动作 / 外部条件"),("同预算","各30k：80%工业 + 20%SCM"),("同协议","960 / 256，full-state，hr0")]):
    x=.83+i*4.23; text(s,t,x,5.06,3.84,.42,23,"teal",True); text(s,d,x,5.64,3.86,.54,17,"muted")
takeaway(s,"两个微调模型冻结ICL predictor；CAQ另训练动作模块。未微调 ≠ 没有context。")

# 11
s=slide("指标口径：不同物理单位，先归一化再聚合", "主指标：固定物理参考尺度归一化MSE ↓，不是原始物理单位MSE", "每个目标误差除预先规定参考尺度r，再平方对query平均；目标等权平均，再设备等权平均。参考尺度与query答案无关且本轮未改变。support标准化是训练编码，与评测参考尺度不同。物理单位RMSE/MAE另保存在JSON。")
steps=[("物理误差","预测 − 真值","K / J / m…"),("固定尺度","误差 ÷ r","预先规定"),("每目标MSE","平方 → query平均","256条预测"),("宏平均","目标 → 设备","52台设备")]
for i,(a,b_,d) in enumerate(steps):
    x=.75+i*3.21; box(s,x,2.24,2.68,2.63,"light","border")
    text(s,a,x+.17,2.46,2.31,.46,23,"teal",True); text(s,b_,x+.17,3.24,2.30,.83,22)
    text(s,d,x+.18,4.26,2.31,.33,16,"muted")
    if i<3: line(s,x+2.75,3.53,x+3.11,3.53,arrow=True)
text(s,"context标准化",1.10,5.43,4.8,.47,24,"blue",True); text(s,"让训练输入/标签数值合适",1.10,5.97,4.9,.38,18,"muted")
text(s,"固定参考尺度",7.34,5.43,4.8,.47,24,"teal",True); text(s,"让不同单位的评估误差可聚合",7.34,5.97,4.8,.38,18,"muted")
takeaway(s,"三模型使用相同任务、目标和参考尺度；没有通过修改尺度降低MSE。")

# 12
s=slide("训练结果：CAQ后期更稳，但最佳不在最后一步", "seed42｜30k成功更新｜无跳步 / 无回滚｜验证仅用于选择best.pt", "曲线来自evaluations.json，从1k开始，不把初始化分数当官方fit/predict基线。CAQ最佳0.10612@19k，最后0.118915；FT最佳0.206764@15k，最后0.311593。验证最佳比较CAQ低48.7%，该数字不能代替测试结果。共享GPU1条件下实际训练和评估耗时约3h58/3h50，不能做独占效率结论。", "来源：industrial-v2/*/evaluations.json 与 summary.json")
picture(s,ASSETS/"validation.png",.68,1.69,8.76,4.54)
for i,(name,k,color) in enumerate([("CAQ","caq","teal"),("TabICL-FT","tabicl","blue")]):
    y=1.93+i*2.03; r=validation["runs"][k]; box(s,9.70,y,2.91,1.73,"light")
    text(s,name,9.89,y+.14,2.51,.39,22,color,True)
    text(s,f"best  {r['best_mse']:.4f}",9.89,y+.69,2.49,.47,25,color,True)
    text(s,f"{r['best_step']//1000}k步  /  final {r['final_mse']:.4f}",9.89,y+1.31,2.52,.27,13,"muted")
takeaway(s,"正式测试提前锁定CAQ 19k与FT 15k；不能用测试集再挑更好的checkpoint。")

# 13
s=slide("正式测试：CAQ较同预算FT误差低10.7%", "52个留出设备｜checkpoint已锁定｜同一typed输入协议｜单训练seed", "三模型meta_test主指标分别0.377522859、0.422901573、9.438638782。官方8成员保留默认预处理、过滤、集成，使用新版单位编码，不是旧frozen core。左图log轴，右图线性轴。验证优势48.7%缩至测试10.7%，只支持本协议单seed平均优势。", "来源：meta-test-seed42-context960-query256-official8/report.json")
picture(s,ASSETS/"test_comparison.png",.63,1.67,9.41,4.61)
stat(s,"10.7%","CAQ相对FT误差降低",10.22,2.02,2.43)
stat(s,"96.0%","CAQ相对官方误差降低",10.22,3.68,2.43)
takeaway(s,"工业训练收益明显；CAQ额外结构优势较小。官方基线结论仅限当前输入协议。")

# 14
s=slide("按原型看：有优势，也有仍然落后的场景", "正值：CAQ更好；负值：FT更好。均使用同一个锁定best.pt，不逐场景挑模型", "8类原型中CAQ胜6类，但仅24/52设备获胜。耦合液罐和换热器分别比FT高20.0%与11.6%。不应只用平均优势宣称普遍优势。辅助原型等权MSE为CAQ0.24728/FT0.28589，低13.5%，不替换主指标。", "来源：正式meta_test报告的scenario_macro_fixed_scale_mse")
picture(s,ASSETS/"scenario_reductions.png",.61,1.64,8.18,4.66)
stat(s,"6 / 8","原型平均误差更低",9.27,1.99,3.13)
stat(s,"24 / 52","逐设备误差更低",9.27,3.52,3.13,"blue")
text(s,"落后：耦合液罐、换热器",9.26,5.15,3.45,.80,19,"orange",True)
takeaway(s,"不同设备、不同机制上的收益不均匀；保留失败场景比只报总分更有解释力。")

# 15
s=slide("目前短板与下一步：维护重置，而不是继续盲目加量", "world444（三工位串联）约占两个微调模型总误差的61%；难例未删除", "设备444 station_0_wear参考尺度固定0.001。context maintenance_0=1有64行，负磨损增量24行；query分别13/5行。context最大磨损下降约0.1484，query约0.3576。不等于完全无维护示例，是更大幅度外推困难的证据，尚不是全部失败原因。未来调整需与已完成正式测试分开，不能反向调测试集。")
for x,v,label,color in [(1.18,.1484,"context中最大下降","blue"),(5.43,.3576,"query中最大下降","orange")]:
    box(s,x,2.09,3.83,2.28,"light","border")
    text(s,label,x+.18,2.30,3.45,.42,21,color,True)
    text(s,f"{v:.3f}",x+.23,3.05,3.36,.69,40,color,True)
    text(s,"磨损增量的绝对幅度",x+.22,3.89,3.37,.30,14,"muted")
box(s,9.72,2.09,2.74,2.28,"navy"); text(s,"不是零维护",9.92,2.36,2.35,.51,23,"white",True)
text(s,"context含维护示例\n但重置范围更小",9.93,3.20,2.33,.94,18,"white")
for i,(t,d) in enumerate([("① 查覆盖","在训练/验证集诊断维护幅度"),("② 做消融","同预算检验动作通道贡献"),("③ 补泛化","多seed与独立机制/场景测试")]):
    x=.80+i*4.20; text(s,t,x,5.08,3.87,.42,22,"teal",True); text(s,d,x,5.70,3.83,.44,16,"muted")
takeaway(s,"当前是单seed模拟设备测试；不是未见机制迁移，更不是已验证真实工厂可用。")

# 16 backup
s=slide("备份：正式测试原型结果与核对入口", "所有数值均为参考尺度归一化MSE；本页按需展示", "主稿件与数据摘要、训练/测试原始JSON为核对依据。原稿引用Modelica、PC-Gym、NIST作为场景组成参考，并不沿用它们的设备精度结论。总结口径：多机制工业合成任务先验+CAQ专用动作通道，单seed未见设备平均收益，需要机制消融与独立迁移进一步证实。")
columns=[.82,4.05,6.45,8.87,11.39]
widths=[2.99,2.15,2.15,2.20,.67]
head=["原型","CAQ","TabICL-FT","官方未微调","设备"]
box(s,.75,1.65,11.89,.52,"navy",rounded=False)
for x,w,t in zip(columns,widths,head): text(s,t,x,1.75,w,.34,17,"white",True)
for i,(key,label) in enumerate(zip(order,cn)):
    y=2.19+i*.43
    if i%2==0: box(s,.75,y,11.89,.42,"light",rounded=False)
    vals=[label,*[f"{test['models'][m]['prediction']['scenario_macro_fixed_scale_mse'][key]:.6g}" for m in ("caq","tabicl_ft","tabicl_official_zero_shot")],str(ca["scenario_world_counts"][key])]
    for j,(x,w,t) in enumerate(zip(columns,widths,vals)): text(s,t,x,y+.065,w,.31,16,"teal" if j==1 else "ink",j==1)
text(s,"来源",.83,5.94,1.15,.34,18,"teal",True)
text(s,"方法稿件 + dataset-summary.json\n训练：formal-comparison.json；测试：report.json / analysis.md",2.02,5.89,10.33,.68,15,"muted")
takeaway(s,"详细绝对路径与补充口径已放在各页演讲者备注，避免主页面塞满文件名。")

outline[0] = ("ICDS工业合成数据与CAQ-TabICL预测实验", outline[0][1])
out = EXP / "ICDS与CAQ-TabICL_组会汇报_20261012.pptx"
prs.save(out)
# Structural geometry check: all objects are inside the slide canvas.
for i,s in enumerate(prs.slides,1):
    for sh in s.shapes:
        assert min(sh.left,sh.top) >= -Inches(.02), (i, sh.name, "negative position")
        assert sh.left+sh.width <= prs.slide_width+Inches(.03), (i,sh.name,"right overflow")
        assert sh.top+sh.height <= prs.slide_height+Inches(.03), (i,sh.name,"bottom overflow")
assert len(prs.slides)==16
with zipfile.ZipFile(out) as z:
    assert len([n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]) == 16
Path(__file__).with_name("汇报提纲.md").write_text("# 组会汇报提纲（15页主报告 + 1页备份）\n\n" +
    "\n\n".join(f"## {i:02d} {title}\n\n{note}" for i,(title,note) in enumerate(outline,1)) +
    f"\n\n来源：{DOC}\n{DATA/'dataset-summary.json'}\n{RESULTS/'formal-comparison.json'}\n{TEST/'report.json'}\n")
print(json.dumps({"pptx": str(out), "slides": len(prs.slides), "bytes": out.stat().st_size,
                  "test_scores": [test['models'][m]['prediction']['macro_fixed_scale_mse'] for m in ("caq","tabicl_ft","tabicl_official_zero_shot")]},ensure_ascii=False))
